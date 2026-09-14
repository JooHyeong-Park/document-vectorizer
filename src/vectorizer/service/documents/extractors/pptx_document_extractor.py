import io


class PptxDocumentExtractor:
    def extract(self, content: bytes) -> str:
        from pptx import Presentation

        return "\n".join(
            shape.text
            for slide in Presentation(io.BytesIO(content)).slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
