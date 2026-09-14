import io


class XlsxDocumentExtractor:
    def extract(self, content: bytes) -> str:
        from openpyxl import load_workbook

        workbook = None
        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            return "\n".join(
                " | ".join("" if value is None else str(value) for value in row)
                for sheet in workbook
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            if workbook is not None:
                workbook.close()
